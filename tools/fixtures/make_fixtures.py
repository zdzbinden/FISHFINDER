# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx", "openpyxl", "pymupdf"]
# ///
"""Regenerate the tools/csp-smoke.js fixtures: one small file per LOAD format.

    uv run tools/fixtures/make_fixtures.py

Each file carries a sentinel string next to a real binomial, so the smoke check
can prove that the matching third-party reader (mammoth, SheetJS, pdf.js) both
loaded under the page's CSP and actually extracted the text. The .xlsx has two
sheets because FISHFINDER reads every sheet, not just the first.

The outputs are committed; the Office and PDF formats embed timestamps, so a
re-run changes bytes without changing content. Nothing here is deployed:
tools/ sits outside fishfinder/.
"""
from pathlib import Path

import docx
import openpyxl
import pymupdf

HERE = Path(__file__).resolve().parent

SENTINELS = {
    "txt": "FIXTURE-TXT Micropterus nigricans",
    "docx": "FIXTURE-DOCX Micropterus nigricans",
    "xlsx": ("FIXTURE-XLSX-SHEET1 Micropterus nigricans",
             "FIXTURE-XLSX-SHEET2 Lepomis macrochirus"),
    "pdf": "FIXTURE-PDF Micropterus nigricans",
}


def main():
    (HERE / "sample.txt").write_text(SENTINELS["txt"] + "\n", encoding="utf-8")

    d = docx.Document()
    d.add_paragraph(SENTINELS["docx"])
    d.save(HERE / "sample.docx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = SENTINELS["xlsx"][0]
    wb.create_sheet("Sheet2")["A1"] = SENTINELS["xlsx"][1]
    wb.save(HERE / "sample.xlsx")

    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), SENTINELS["pdf"], fontsize=12)
    pdf.save(HERE / "sample.pdf")

    for f in sorted(HERE.glob("sample.*")):
        print(f"{f.name:12s} {f.stat().st_size:6d} bytes")


if __name__ == "__main__":
    main()
